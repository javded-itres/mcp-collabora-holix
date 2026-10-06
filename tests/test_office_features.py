"""Fonts, media, slides, formulas, charts, and tables in the office MCP."""

from __future__ import annotations

import json
import zipfile
from io import BytesIO
from pathlib import Path

from docx import Document
from odf.opendocument import OpenDocumentPresentation, OpenDocumentSpreadsheet, OpenDocumentText
from openpyxl import Workbook, load_workbook
from PIL import Image
from pptx import Presentation

from holix_office.charts import chart_office
from holix_office.edit import format_office_text
from holix_office.media import insert_office_media
from holix_office.odf_xml import ods_formula
from holix_office.sheets import sheet_office
from holix_office.slides import slide_office
from holix_office.tables import table_office


def _png(path: Path) -> None:
    handle = BytesIO()
    Image.new("RGB", (8, 8), (200, 20, 20)).save(handle, format="PNG")
    path.write_bytes(handle.getvalue())


def _mp4(path: Path) -> None:
    path.write_bytes(b"\x00\x00\x00\x18ftypmp42\x00\x00\x00\x00mp42isom")


def test_font_on_docx_pptx_xlsx_and_odt(tmp_path: Path) -> None:
    docx = tmp_path / "note.docx"
    doc = Document()
    doc.add_paragraph("Hello Holix")
    doc.save(docx)
    assert format_office_text(docx, "Holix", font="Calibri", size=18, underline=True) == 1
    xml = zipfile.ZipFile(docx).read("word/document.xml").decode()
    assert 'w:ascii="Calibri"' in xml
    assert 'w:val="36"' in xml

    deck = tmp_path / "deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Hello Slide"
    prs.save(deck)
    assert format_office_text(deck, "Slide", bold=True, font="Calibri", size=20, color="FF0000") == 1
    slide_xml = zipfile.ZipFile(deck).read("ppt/slides/slide1.xml").decode()
    assert 'typeface="Calibri"' in slide_xml
    assert 'sz="2000"' in slide_xml

    book = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active["A1"] = "Total"
    wb.save(book)
    assert format_office_text(book, "Total", font="Calibri", size=16, bold=True) == 1
    assert load_workbook(book).active["A1"].font.name == "Calibri"
    assert load_workbook(book).active["A1"].font.size == 16

    odt = tmp_path / "note.odt"
    OpenDocumentText().save(odt)
    # Blank ODF has no paragraph yet; append one through the text body via a tiny package edit.
    from holix_office.edit import append_office_paragraph

    # odfpy saves an empty text body, and append supports odt.
    append_office_paragraph(odt, "Hello Writer")
    assert format_office_text(odt, "Writer", font="Liberation Serif", size=14) == 1
    content = zipfile.ZipFile(odt).read("content.xml").decode()
    assert "Liberation Serif" in content
    assert "14pt" in content


def test_named_heading_style(tmp_path: Path) -> None:
    path = tmp_path / "note.docx"
    doc = Document()
    doc.add_paragraph("Title")
    doc.save(path)
    assert format_office_text(path, "Title", style="Heading 1") == 1
    assert Document(str(path)).paragraphs[0].style.name == "Heading 1"


def test_formula_link_chart_and_table_xlsx(tmp_path: Path) -> None:
    path = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active.title = "Data"
    wb.active["A1"] = "Name"
    wb.active["B1"] = "Qty"
    wb.save(path)
    added = sheet_office(path, "add", sheet_name="Other")
    assert added["sheet"] == "Other"
    sheet_office(path, "set", sheet_name="Other", cell="A1", value="7")
    linked = sheet_office(path, "formula", sheet_name="Data", cell="B2", formula="=Other!A1")
    assert linked["value"] == "=Other!A1"
    sheet_office(path, "link", sheet_name="Data", cell="A2", value="Site", link="https://example.com")
    named = sheet_office(path, "name", sheet_name="Qty", range_ref="Data!$B$1:$B$2")
    assert named["ref"] == "Data!$B$1:$B$2"
    chart_office(
        path,
        chart_type="bar",
        sheet_name="Data",
        data_range="B1:B2",
        categories_range="A2:A2",
        anchor="E2",
        title="Qty",
    )
    book = load_workbook(path)
    assert book["Data"]["B2"].value == "=Other!A1"
    assert book["Data"]["A2"].hyperlink.target == "https://example.com"
    assert any("chart" in name for name in zipfile.ZipFile(path).namelist())
    table_office(path, "add", rows=[["Name", "Qty"], ["A", "1"]])
    book = load_workbook(path)
    assert "Table1" in book["Data"].tables
    listed = sheet_office(path, "read", sheet_name="Other")
    assert any(item["cell"] == "A1" and str(item["value"]) in {"7", "7.0"} for item in listed["cells"])


def test_ods_formula_and_slide_media(tmp_path: Path) -> None:
    assert ods_formula("=Other!A1") == "of:=[Other.A1]"
    assert ods_formula("=SUM(A1:A2)") == "of:=SUM([.A1:.A2])"
    ods = tmp_path / "book.ods"
    OpenDocumentSpreadsheet().save(ods)
    sheet_office(ods, "add", sheet_name="Sheet1")
    sheet_office(ods, "add", sheet_name="Other")
    sheet_office(ods, "set", sheet_name="Other", cell="A1", value="7")
    sheet_office(ods, "formula", sheet_name="Sheet1", cell="B1", formula="=Other!A1")
    content = zipfile.ZipFile(ods).read("content.xml").decode()
    assert "of:=[Other.A1]" in content

    png = tmp_path / "cover.png"
    _png(png)
    odp = tmp_path / "deck.odp"
    OpenDocumentPresentation().save(odp)
    added = slide_office(odp, "add", title="Hello", body_text="World")
    assert added["action"] == "add"
    listed = slide_office(odp, "list")
    assert listed["slides"][0]["text"] == ["Hello", "World"]
    insert_office_media(odp, png, slide=1, width_cm=8, height_cm=4)
    names = zipfile.ZipFile(odp).namelist()
    assert "Pictures/cover.png" in names
    assert "image/png" in zipfile.ZipFile(odp).read("META-INF/manifest.xml").decode()

    video = tmp_path / "clip.mp4"
    _mp4(video)
    deck = tmp_path / "deck.pptx"
    Presentation().save(deck)
    slide_office(deck, "add", title="Talk", body_text="Body", layout="title_body")
    slide_office(deck, "notes", slide=1, notes="Say this")
    insert_office_media(deck, video, slide=1)
    packed = zipfile.ZipFile(deck).namelist()
    assert any(name.endswith(".mp4") for name in packed)
    assert Presentation(str(deck)).slides[0].has_notes_slide

    docx = tmp_path / "note.docx"
    Document().save(docx)
    insert_office_media(docx, png, width_cm=4)
    assert any("media" in name for name in zipfile.ZipFile(docx).namelist())
    table_office(docx, "add", rows=json.loads('[["Name","Qty"],["A","1"]]'))
    saved = Document(str(docx))
    assert saved.tables[0].rows[1].cells[1].text == "1"
    table_office(docx, "set", table=1, row=1, col=1, value="9")
    assert Document(str(docx)).tables[0].rows[1].cells[1].text == "9"


def test_pptx_chart(tmp_path: Path) -> None:
    path = tmp_path / "deck.pptx"
    Presentation().save(path)
    slide_office(path, "add", layout="blank")
    chart_office(
        path,
        chart_type="line",
        slide=1,
        title="Trend",
        categories="Jan,Feb",
        series_name="Sales",
        values="10,12",
    )
    assert any("chart" in name for name in zipfile.ZipFile(path).namelist())
