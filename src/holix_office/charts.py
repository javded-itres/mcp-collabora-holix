"""Charts bound to a worksheet range, or drawn on a slide from a series."""

from __future__ import annotations

from pathlib import Path

from holix_office.edit import OfficeEditError, _require_office_file

_XLSX = {
    "bar": "bar",
    "column": "bar",
    "line": "line",
    "pie": "pie",
}
_PPTX = {
    "bar": "COLUMN_CLUSTERED",
    "column": "COLUMN_CLUSTERED",
    "line": "LINE",
    "pie": "PIE",
}


def chart_office(
    path: Path,
    *,
    chart_type: str = "bar",
    sheet_name: str = "",
    data_range: str = "",
    categories_range: str = "",
    anchor: str = "E2",
    title: str = "",
    slide: int = 1,
    categories: str = "",
    series_name: str = "",
    values: str = "",
) -> dict[str, object]:
    path = _require_office_file(path)
    kind = (chart_type or "bar").strip().lower()
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        if kind not in _XLSX:
            raise OfficeEditError("chart_type must be bar, line, or pie")
        _xlsx_chart(
            path,
            kind=kind,
            sheet_name=sheet_name,
            data_range=data_range,
            categories_range=categories_range,
            anchor=anchor or "E2",
            title=title,
        )
    elif suffix == ".pptx":
        if kind not in _PPTX:
            raise OfficeEditError("chart_type must be bar, line, or pie")
        _pptx_chart(
            path,
            kind=kind,
            slide=slide,
            title=title,
            categories=categories,
            series_name=series_name,
            values=values,
        )
    else:
        raise OfficeEditError("Charts are added to xlsx worksheets and pptx slides")
    return {"ok": True, "chart_type": kind, "path": path.name}


def _xlsx_chart(path, *, kind, sheet_name, data_range, categories_range, anchor, title) -> None:
    from openpyxl import load_workbook
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference

    if not data_range.strip():
        raise OfficeEditError("data_range is required, for example B1:B4")
    book = load_workbook(path)
    if sheet_name.strip():
        if sheet_name.strip() not in book.sheetnames:
            raise OfficeEditError(f"Sheet {sheet_name} was not found")
        worksheet = book[sheet_name.strip()]
    else:
        worksheet = book.active
    chart = {"bar": BarChart, "line": LineChart, "pie": PieChart}[kind]()
    if title:
        chart.title = title
    data = Reference(worksheet, range_string=f"'{worksheet.title}'!{data_range.strip()}")
    chart.add_data(data, titles_from_data=True)
    if categories_range.strip():
        chart.set_categories(
            Reference(worksheet, range_string=f"'{worksheet.title}'!{categories_range.strip()}")
        )
    if kind == "bar":
        chart.shape = 4
    worksheet.add_chart(chart, (anchor or "E2").strip().upper())
    book.save(path)


def _pptx_chart(path, *, kind, slide, title, categories, series_name, values) -> None:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Cm

    labels = [item.strip() for item in categories.split(",") if item.strip()]
    numbers = []
    for item in values.split(","):
        text = item.strip()
        if not text:
            continue
        try:
            numbers.append(float(text) if "." in text else int(text))
        except ValueError as exc:
            raise OfficeEditError("values must be comma-separated numbers") from exc
    if not labels or not numbers or len(labels) != len(numbers):
        raise OfficeEditError("categories and values must be comma-separated lists of the same length")
    deck = Presentation(str(path))
    if len(deck.slides) == 0:
        target = deck.slides.add_slide(deck.slide_layouts[6])
    else:
        index = max(1, int(slide or 1)) - 1
        if index >= len(deck.slides):
            raise OfficeEditError("Slide was not found")
        target = deck.slides[index]
    data = CategoryChartData()
    data.categories = labels
    data.add_series(series_name.strip() or "Series 1", numbers)
    graphic = target.shapes.add_chart(
        getattr(XL_CHART_TYPE, _PPTX[kind]),
        Cm(1),
        Cm(3),
        Cm(16),
        Cm(9),
        data,
    )
    if title:
        graphic.chart.has_title = True
        graphic.chart.chart_title.text_frame.text = title
    deck.save(str(path))
