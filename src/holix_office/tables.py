"""Document tables, slide tables, and Excel tables."""

from __future__ import annotations

import json
import xml.etree.ElementTree as ET
from pathlib import Path

from holix_office.edit import OfficeEditError, _require_office_file
from holix_office.odf_xml import (
    DRAW,
    TABLE,
    TEXT,
    body,
    edit_content,
    pages,
    read_content,
)

_MAX_ROWS = 200
_MAX_COLS = 40


def table_office(
    path: Path,
    action: str,
    *,
    rows: list[list[str]] | None = None,
    table: int = 1,
    row: int = 0,
    col: int = 0,
    value: str = "",
    slide: int = 1,
    header: bool = True,
) -> dict[str, object]:
    path = _require_office_file(path)
    verb = (action or "").strip().lower()
    suffix = path.suffix.lower()
    if verb == "add":
        grid = _grid(rows)
        _add(path, suffix, grid, slide=slide, header=header)
        return {"ok": True, "action": "add", "rows": len(grid), "cols": len(grid[0])}
    if verb == "read":
        return {"ok": True, "tables": _read(path, suffix)}
    if verb == "set":
        _set(path, suffix, table=table, row=row, col=col, value=value, slide=slide)
        return {"ok": True, "action": "set", "table": table, "row": row, "col": col}
    raise OfficeEditError("action must be read, add, or set")


def parse_rows(payload: str) -> list[list[str]]:
    if not str(payload or "").strip():
        raise OfficeEditError("rows_json is empty")
    try:
        raw = json.loads(payload)
    except json.JSONDecodeError as exc:
        raise OfficeEditError("rows_json must be a JSON list of rows") from exc
    if not isinstance(raw, list) or not raw:
        raise OfficeEditError("rows_json must be a JSON list of rows")
    rows = []
    width = 0
    for line in raw:
        if not isinstance(line, list):
            raise OfficeEditError("each row must be a JSON list")
        rows.append(["" if item is None else str(item) for item in line])
        width = max(width, len(line))
    if width == 0:
        raise OfficeEditError("table has no columns")
    return [line + [""] * (width - len(line)) for line in rows]


def _grid(rows: list[list[str]] | None) -> list[list[str]]:
    grid = rows or []
    if not grid or not grid[0]:
        raise OfficeEditError("table has no cells")
    if len(grid) > _MAX_ROWS or len(grid[0]) > _MAX_COLS:
        raise OfficeEditError("table is too large")
    return grid


def _add(path: Path, suffix: str, grid: list[list[str]], *, slide: int, header: bool) -> None:
    if suffix == ".docx":
        _docx_add(path, grid)
    elif suffix == ".pptx":
        _pptx_add(path, grid, slide)
    elif suffix == ".xlsx":
        _xlsx_add(path, grid, header)
    elif suffix in {".odt", ".odp"}:
        _odf_add(path, grid, slide if suffix == ".odp" else 0)
    elif suffix == ".ods":
        _ods_grid(path, grid)
    else:
        raise OfficeEditError("Tables can be added to docx, xlsx, pptx, odt, ods, and odp")


def _read(path: Path, suffix: str) -> list[dict[str, object]]:
    if suffix == ".docx":
        from docx import Document

        doc = Document(str(path))
        return [_docx_table(item, index) for index, item in enumerate(doc.tables, start=1)]
    if suffix == ".pptx":
        from pptx import Presentation

        deck = Presentation(str(path))
        found = []
        number = 1
        for slide_index, slide in enumerate(deck.slides, start=1):
            for shape in slide.shapes:
                if not getattr(shape, "has_table", False):
                    continue
                found.append({"table": number, "slide": slide_index, "rows": _pptx_rows(shape.table)})
                number += 1
        return found
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        book = load_workbook(path)
        found = []
        for worksheet in book.worksheets:
            for item in worksheet.tables.values():
                found.append({"name": item.name, "sheet": worksheet.title, "ref": item.ref})
        return found
    if suffix in {".odt", ".ods", ".odp"}:
        root = read_content(path)
        found = []
        for index, table in enumerate(root.iter(), start=1):
            if table.tag != f"{{{TABLE}}}table":
                continue
            if suffix == ".ods":
                continue
            found.append({"table": len(found) + 1, "rows": _odf_rows(table)})
        return found
    raise OfficeEditError("Tables can be read from docx, xlsx, pptx, odt, ods, and odp")


def _set(path, suffix, *, table, row, col, value, slide) -> None:
    if row < 0 or col < 0:
        raise OfficeEditError("row and col start at 0")
    if suffix == ".docx":
        from docx import Document

        doc = Document(str(path))
        chosen = _at(doc.tables, table, "table")
        chosen.rows[row].cells[col].text = value
        doc.save(str(path))
        return
    if suffix == ".pptx":
        from pptx import Presentation

        deck = Presentation(str(path))
        tables = [
            shape.table
            for item in deck.slides
            for shape in item.shapes
            if getattr(shape, "has_table", False)
        ]
        chosen = _at(tables, table, "table")
        chosen.cell(row, col).text = value
        deck.save(str(path))
        return
    if suffix in {".odt", ".odp"}:
        def change(root: ET.Element) -> None:
            tables = [el for el in root.iter() if el.tag == f"{{{TABLE}}}table"]
            if suffix == ".odp" and slide:
                page = _at(pages(root), slide, "slide")
                tables = [el for el in page.iter() if el.tag == f"{{{TABLE}}}table"] or tables
            chosen = _at(tables, table, "table")
            rows = [el for el in list(chosen) if el.tag == f"{{{TABLE}}}table-row"]
            cells = [
                el
                for el in list(rows[row])
                if el.tag == f"{{{TABLE}}}table-cell"
            ]
            paragraph = cells[col].find(f"{{{TEXT}}}p")
            if paragraph is None:
                paragraph = ET.SubElement(cells[col], f"{{{TEXT}}}p")
            paragraph.text = value
            for child in list(paragraph):
                paragraph.remove(child)

        edit_content(path, change)
        return
    raise OfficeEditError("Setting one cell is available for docx, pptx, odt, and odp tables")


def _docx_add(path: Path, grid: list[list[str]]) -> None:
    from docx import Document

    doc = Document(str(path))
    table = doc.add_table(rows=len(grid), cols=len(grid[0]))
    table.style = "Table Grid"
    for r, line in enumerate(grid):
        for c, item in enumerate(line):
            table.rows[r].cells[c].text = item
    doc.save(str(path))


def _pptx_add(path: Path, grid: list[list[str]], slide: int) -> None:
    from pptx import Presentation
    from pptx.util import Cm

    deck = Presentation(str(path))
    if len(deck.slides) == 0:
        target = deck.slides.add_slide(deck.slide_layouts[6])
    else:
        target = deck.slides[min(max(1, slide) - 1, len(deck.slides) - 1)]
    shape = target.shapes.add_table(len(grid), len(grid[0]), Cm(1), Cm(3), Cm(20), Cm(10))
    for r, line in enumerate(grid):
        for c, item in enumerate(line):
            shape.table.cell(r, c).text = item
    deck.save(str(path))


def _xlsx_add(path: Path, grid: list[list[str]], header: bool) -> None:
    from openpyxl import load_workbook
    from openpyxl.worksheet.table import Table, TableStyleInfo

    book = load_workbook(path)
    worksheet = book.active
    for r, line in enumerate(grid, start=1):
        for c, item in enumerate(line, start=1):
            worksheet.cell(r, c).value = item
    ref = f"A1:{_col(len(grid[0]))}{len(grid)}"
    name = f"Table{len(worksheet.tables) + 1}"
    item = Table(displayName=name, ref=ref)
    item.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    if header:
        item.headerRowCount = 1
    worksheet.add_table(item)
    book.save(path)


def _odf_add(path: Path, grid: list[list[str]], slide: int) -> None:
    def change(root: ET.Element) -> None:
        table = _odf_table(grid)
        if slide:
            found = pages(root)
            if not found:
                page = ET.SubElement(body(root, "presentation"), f"{{{DRAW}}}page")
                found = [page]
            found[min(slide - 1, len(found) - 1)].append(table)
        else:
            body(root, "text").append(table)

    edit_content(path, change)


def _ods_grid(path: Path, grid: list[list[str]]) -> None:
    from holix_office.odf_xml import put_cell, sheet

    def change(root: ET.Element) -> None:
        table = sheet(root, "")
        for r, line in enumerate(grid):
            for c, item in enumerate(line):
                put_cell(table, f"{_col(c + 1)}{r + 1}", item)

    edit_content(path, change)


def _odf_table(grid: list[list[str]]) -> ET.Element:
    table = ET.Element(f"{{{TABLE}}}table")
    for line in grid:
        row = ET.SubElement(table, f"{{{TABLE}}}table-row")
        for item in line:
            cell = ET.SubElement(row, f"{{{TABLE}}}table-cell")
            paragraph = ET.SubElement(cell, f"{{{TEXT}}}p")
            paragraph.text = item
    return table


def _docx_table(table, index: int) -> dict[str, object]:
    return {
        "table": index,
        "rows": [[cell.text for cell in row.cells] for row in table.rows],
    }


def _pptx_rows(table) -> list[list[str]]:
    return [[cell.text for cell in row.cells] for row in table.rows]


def _odf_rows(table: ET.Element) -> list[list[str]]:
    rows = []
    for row in list(table):
        if row.tag != f"{{{TABLE}}}table-row":
            continue
        rows.append(
            [
                "".join(cell.itertext())
                for cell in list(row)
                if cell.tag == f"{{{TABLE}}}table-cell"
            ]
        )
    return rows


def _at(items, number: int, label: str):
    index = max(1, int(number or 1)) - 1
    if index >= len(items):
        raise OfficeEditError(f"{label.capitalize()} was not found")
    return items[index]


def _col(number: int) -> str:
    label = ""
    while number:
        number, rem = divmod(number - 1, 26)
        label = chr(65 + rem) + label
    return label
