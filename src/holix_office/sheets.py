"""Workbook sheets, formulas, defined names, and cell links."""

from __future__ import annotations

from pathlib import Path

from holix_office.edit import OfficeEditError, _require_office_file
from holix_office.odf_xml import (
    add_sheet,
    edit_content,
    put_cell,
    read_content,
    read_sheets,
    sheet,
)


def sheet_office(
    path: Path,
    action: str,
    *,
    sheet_name: str = "",
    cell: str = "",
    value: str = "",
    formula: str = "",
    link: str = "",
    range_ref: str = "",
) -> dict[str, object]:
    path = _require_office_file(path)
    suffix = path.suffix.lower()
    if suffix not in {".xlsx", ".ods"}:
        raise OfficeEditError("Sheets, formulas, and links are edited in xlsx and ods files")
    verb = (action or "").strip().lower()
    if suffix == ".xlsx":
        return _xlsx(
            path,
            verb,
            sheet_name=sheet_name,
            cell=cell,
            value=value,
            formula=formula,
            link=link,
            range_ref=range_ref,
        )
    return _ods(
        path,
        verb,
        sheet_name=sheet_name,
        cell=cell,
        value=value,
        formula=formula,
        link=link,
        range_ref=range_ref,
    )


def _xlsx(path, action, *, sheet_name, cell, value, formula, link, range_ref) -> dict[str, object]:
    from openpyxl import load_workbook
    from openpyxl.workbook.defined_name import DefinedName

    book = load_workbook(path)
    if action == "list":
        return {
            "ok": True,
            "sheets": [_sheet_payload(item) for item in book.worksheets],
            "names": [
                {"name": defined.name, "ref": defined.attr_text}
                for defined in book.defined_names.values()
            ],
        }
    if action == "add":
        name = _sheet_title(sheet_name, book.sheetnames)
        if name in book.sheetnames:
            raise OfficeEditError(f"Sheet {name} already exists")
        book.create_sheet(name)
        book.save(path)
        return {"ok": True, "action": "add", "sheet": name}
    if action == "name":
        if not sheet_name.strip() or not range_ref.strip():
            raise OfficeEditError("A defined name needs a name and a range such as Sheet1!$A$1:$B$10")
        attr = range_ref.strip()
        if sheet_name in book.defined_names:
            del book.defined_names[sheet_name]
        book.defined_names.add(DefinedName(name=sheet_name.strip(), attr_text=attr))
        book.save(path)
        return {"ok": True, "action": "name", "name": sheet_name.strip(), "ref": attr}
    target = _xlsx_sheet(book, sheet_name)
    if action == "read":
        return {"ok": True, "sheet": target.title, "cells": _sheet_payload(target)["cells"]}
    if action not in {"set", "formula", "link"}:
        raise OfficeEditError("action must be list, add, read, set, formula, link, or name")
    if not cell.strip():
        raise OfficeEditError("Cell must look like A1")
    chosen = target[cell.strip().upper()]
    if action == "formula" or formula.strip():
        text = formula.strip() or value.strip()
        if not text:
            raise OfficeEditError("Formula is empty")
        if not text.startswith("="):
            text = "=" + text
        chosen.value = text
    elif action == "set":
        chosen.value = _coerce(value)
    if link.strip():
        chosen.hyperlink = link.strip()
        if chosen.value in (None, ""):
            chosen.value = link.strip()
    book.save(path)
    return {
        "ok": True,
        "action": action,
        "sheet": target.title,
        "cell": chosen.coordinate,
        "value": chosen.value,
        "link": link.strip(),
    }


def _ods(path, action, *, sheet_name, cell, value, formula, link, range_ref) -> dict[str, object]:
    if action == "name":
        raise OfficeEditError("Defined names are available for xlsx")
    if action == "list":
        return {"ok": True, "sheets": read_sheets(read_content(path))}
    if action == "add":
        name = sheet_name.strip() or "Sheet"
        edit_content(path, lambda root: add_sheet(root, name))
        return {"ok": True, "action": "add", "sheet": name}
    if action == "read":
        sheets = read_sheets(read_content(path))
        if sheet_name.strip():
            sheets = [item for item in sheets if item["name"] == sheet_name.strip()]
            if not sheets:
                raise OfficeEditError(f"Sheet {sheet_name} was not found")
        return {"ok": True, "sheets": sheets}
    if action not in {"set", "formula", "link"}:
        raise OfficeEditError("action must be list, add, read, set, formula, or link")
    if not cell.strip():
        raise OfficeEditError("Cell must look like A1")

    def change(root) -> None:
        put_cell(
            sheet(root, sheet_name.strip()),
            cell,
            value,
            formula if action == "formula" or formula.strip() else "",
            link,
        )

    edit_content(path, change)
    return {"ok": True, "action": action, "sheet": sheet_name.strip() or "Sheet1", "cell": cell.upper()}


def _sheet_payload(sheet) -> dict[str, object]:
    cells = []
    for row in sheet.iter_rows():
        for item in row:
            if item.value is None and item.hyperlink is None:
                continue
            payload: dict[str, object] = {"cell": item.coordinate, "value": item.value}
            if isinstance(item.value, str) and item.value.startswith("="):
                payload["formula"] = item.value
            if item.hyperlink is not None:
                payload["link"] = item.hyperlink.target
            cells.append(payload)
            if len(cells) >= 400:
                return {"name": sheet.title, "cells": cells}
    return {"name": sheet.title, "cells": cells}


def _xlsx_sheet(book, name: str):
    if not name.strip():
        return book.active
    if name.strip() not in book.sheetnames:
        raise OfficeEditError(f"Sheet {name} was not found")
    return book[name.strip()]


def _sheet_title(name: str, existing: list[str]) -> str:
    title = (name or "").strip() or f"Sheet{len(existing) + 1}"
    if len(title) > 31:
        raise OfficeEditError("Sheet name must be 31 characters or fewer")
    return title


def _coerce(value: str):
    text = "" if value is None else str(value)
    if text == "":
        return ""
    try:
        if "." in text:
            return float(text)
        return int(text)
    except ValueError:
        return text
