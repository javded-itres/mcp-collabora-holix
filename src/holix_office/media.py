"""Insert a workspace image or video into an office document."""

from __future__ import annotations

import re
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from holix_office.edit import (
    OfficeEditError,
    _clone_info,
    _prepare,
    _read_zip,
    _require_office_file,
    _serialize,
    _write_zip,
)
from holix_office.odf_xml import DRAW, body, edit_content, media_frame, pages

_IMAGES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
    ".bmp": "image/bmp",
    ".tif": "image/tiff",
    ".tiff": "image/tiff",
}
_VIDEOS = {
    ".mp4": "video/mp4",
    ".m4v": "video/mp4",
    ".webm": "video/webm",
    ".mov": "video/quicktime",
    ".ogg": "video/ogg",
}
_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
_HYPERLINK = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_CT = "http://schemas.openxmlformats.org/package/2006/content-types"
_MAX_MEDIA = 200_000_000


def insert_office_media(
    path: Path,
    media: Path,
    *,
    slide: int = 1,
    width_cm: float = 12,
    height_cm: float = 6.8,
) -> dict[str, object]:
    path = _require_office_file(path)
    media = media.expanduser().resolve()
    if not media.is_file():
        raise OfficeEditError("Media file not found")
    if media.stat().st_size > _MAX_MEDIA:
        raise OfficeEditError("Media file is too large")
    mime = _IMAGES.get(media.suffix.lower())
    video = False
    if mime is None:
        mime = _VIDEOS.get(media.suffix.lower())
        video = mime is not None
    if mime is None:
        raise OfficeEditError("Media must be png, jpg, gif, webp, bmp, tiff, mp4, webm, or mov")
    suffix = path.suffix.lower()
    if video and suffix in {".xlsx", ".ods"}:
        raise OfficeEditError("Put a video on a slide or in a text document")
    width = _positive(width_cm, "width_cm")
    height = _positive(height_cm, "height_cm")
    index = max(1, int(slide or 1))
    if suffix == ".docx":
        if video:
            _docx_video(path, media, mime)
        else:
            _docx_image(path, media, width)
    elif suffix == ".pptx":
        _pptx_media(path, media, mime, index, width, height, video=video)
    elif suffix == ".xlsx":
        _xlsx_image(path, media, index, width, height)
    elif suffix in {".odt", ".ods", ".odp"}:
        _odf_media(path, media, mime, index, width, height, video=video)
    else:
        raise OfficeEditError("Media can be inserted into docx, xlsx, pptx, odt, ods, and odp")
    return {
        "ok": True,
        "kind": "video" if video else "image",
        "media": media.name,
        "slide": index,
    }


def _docx_image(path: Path, media: Path, width_cm: float) -> None:
    from docx import Document
    from docx.shared import Cm

    doc = Document(str(path))
    doc.add_picture(str(media), width=Cm(width_cm))
    doc.save(str(path))


def _docx_video(path: Path, media: Path, mime: str) -> None:
    safe = _safe_name(media.name)
    target = f"media/{safe}"
    members = _read_zip(path)
    names = {info.filename for info, _data in members}
    updated: list[tuple[zipfile.ZipInfo, bytes]] = []
    rel_id = "rIdHolixVideo"
    for info, data in members:
        if info.filename == "[Content_Types].xml":
            updated.append((_clone_info(info), _content_type(data, media.suffix.lower(), mime)))
        elif info.filename == "word/_rels/document.xml.rels":
            updated.append((_clone_info(info), _relationship(data, rel_id, target)))
        elif info.filename == "word/document.xml":
            updated.append((_clone_info(info), _video_paragraph(data, rel_id, media.name)))
        else:
            updated.append((info, data))
    stored = f"word/{target}"
    if stored not in names:
        info = zipfile.ZipInfo(stored)
        info.compress_type = zipfile.ZIP_DEFLATED
        info.date_time = (2026, 1, 1, 0, 0, 0)
        updated.append((info, media.read_bytes()))
    _write_zip(path, updated)


def _pptx_media(path, media, mime, slide, width, height, *, video: bool) -> None:
    from pptx import Presentation
    from pptx.util import Cm

    deck = Presentation(str(path))
    if len(deck.slides) == 0:
        target = deck.slides.add_slide(deck.slide_layouts[6])
    else:
        target = deck.slides[min(slide - 1, len(deck.slides) - 1)]
    poster = None
    try:
        if video:
            poster = _poster()
            target.shapes.add_movie(
                str(media),
                Cm(1),
                Cm(1),
                Cm(width),
                Cm(height),
                poster_frame_image=str(poster),
                mime_type=mime,
            )
        else:
            target.shapes.add_picture(str(media), Cm(1), Cm(1), Cm(width), Cm(height))
        deck.save(str(path))
    finally:
        if poster is not None:
            poster.unlink(missing_ok=True)


def _xlsx_image(path, media, sheet_index, width, height) -> None:
    from openpyxl import load_workbook
    from openpyxl.drawing.image import Image as XLImage

    book = load_workbook(path)
    sheet = book.worksheets[min(sheet_index - 1, len(book.worksheets) - 1)]
    image = XLImage(str(media))
    image.width = int(width * 37.8)
    image.height = int(height * 37.8)
    sheet.add_image(image, "A1")
    book.save(path)


def _odf_media(path, media, mime, slide, width, height, *, video: bool) -> None:
    folder = "Media" if video else "Pictures"
    href = f"{folder}/{_safe_name(media.name)}"
    blob = media.read_bytes()

    def change(root: ET.Element) -> None:
        frame = media_frame(href, mime, width, height, video=video)
        suffix = path.suffix.lower()
        if suffix == ".odp":
            found = pages(root)
            if not found:
                page = ET.SubElement(body(root, "presentation"), f"{{{DRAW}}}page")
                page.set(f"{{{DRAW}}}name", "page1")
                found = [page]
            found[min(slide - 1, len(found) - 1)].append(frame)
        elif suffix == ".ods":
            body(root, "spreadsheet").append(frame)
        else:
            body(root, "text").append(frame)

    edit_content(path, change, [(href, mime, blob)])


def _poster() -> Path:
    import os
    import tempfile
    from io import BytesIO

    from PIL import Image

    handle = BytesIO()
    Image.new("RGB", (16, 9), (20, 20, 20)).save(handle, format="PNG")
    fd, name = tempfile.mkstemp(suffix=".png")
    os.close(fd)
    path = Path(name)
    path.write_bytes(handle.getvalue())
    return path


def _content_type(payload: bytes, suffix: str, mime: str) -> bytes:
    root = _prepare(payload)
    ext = suffix.lstrip(".")
    for child in root:
        if child.get("Extension") == ext:
            return payload
    node = ET.SubElement(root, f"{{{_CT}}}Default")
    node.set("Extension", ext)
    node.set("ContentType", mime)
    return _serialize(root)


def _relationship(payload: bytes, rel_id: str, target: str) -> bytes:
    root = _prepare(payload)
    for child in root:
        if child.get("Id") == rel_id:
            return payload
    node = ET.SubElement(root, f"{{{_REL}}}Relationship")
    node.set("Id", rel_id)
    node.set("Type", _HYPERLINK)
    node.set("Target", target)
    return _serialize(root)


def _video_paragraph(payload: bytes, rel_id: str, label: str) -> bytes:
    root = _prepare(payload)
    document_body = root.find(f"{{{_W}}}body")
    if document_body is None:
        raise OfficeEditError("Document has no body")
    paragraph = ET.Element(f"{{{_W}}}p")
    link = ET.SubElement(paragraph, f"{{{_W}}}hyperlink")
    link.set(f"{{{_R}}}id", rel_id)
    run = ET.SubElement(link, f"{{{_W}}}r")
    node = ET.SubElement(run, f"{{{_W}}}t")
    node.text = f"Video: {label}"
    children = list(document_body)
    if children and children[-1].tag == f"{{{_W}}}sectPr":
        document_body.insert(len(children) - 1, paragraph)
    else:
        document_body.append(paragraph)
    return _serialize(root)


def _safe_name(name: str) -> str:
    clean = re.sub(r"[^A-Za-z0-9._-]+", "_", Path(name).name).strip("._")
    return clean or "media.bin"


def _positive(value: float, label: str) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise OfficeEditError(f"{label} must be a number of centimetres") from exc
    if number <= 0 or number > 100:
        raise OfficeEditError(f"{label} is out of range")
    return number
