"""Phrase styles for slides and workbooks, plus named paragraph styles."""

from __future__ import annotations

import xml.etree.ElementTree as ET
from pathlib import Path

from holix_office.edit import OfficeEditError

_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_TEXT = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"
_OFFICE = "urn:oasis:names:tc:opendocument:xmlns:office:1.0"
_STYLE = "urn:oasis:names:tc:opendocument:xmlns:style:1.0"
_FO = "urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0"

_HEADING_SIZE = {
    "Title": 28,
    "Heading 1": 18,
    "Heading 2": 16,
    "Heading 3": 14,
}


def format_pptx_tree(
    root: ET.Element,
    needle: str,
    *,
    bold: bool | None,
    italic: bool | None,
    underline: bool | None,
    color: str | None,
    fill: str | None,
    font: str | None,
    size_pt: float | None,
) -> int:
    """Format a phrase on one slide without restyling the words around it."""
    hits = 0
    for para in _drawing_paragraphs(root):
        joined = "".join(node.text or "" for _run, node in _pptx_text_nodes(para))
        shade = bool(fill) and joined.strip() == needle.strip() and bool(needle.strip())
        found = 0
        guard = 0
        while guard < 10000:
            nodes = _pptx_text_nodes(para)
            joined = "".join(node.text or "" for _run, node in nodes)
            start = _mask_right(joined, needle, found).rfind(needle)
            if start < 0:
                break
            _style_span(
                para,
                nodes,
                start,
                start + len(needle),
                bold=bold,
                italic=italic,
                underline=underline,
                color=color,
                font=font,
                size_pt=size_pt,
            )
            found += 1
            guard += 1
        if found and shade:
            _shade_paragraph(para, fill)
        hits += found
    return hits


def format_xlsx_cells(
    path: Path,
    needle: str,
    *,
    bold: bool | None,
    italic: bool | None,
    underline: bool | None,
    color: str | None,
    fill: str | None,
    font: str | None,
    size_pt: float | None,
) -> int:
    """Apply a font to every cell whose value or formula contains the phrase."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill

    try:
        book = load_workbook(path)
    except Exception as exc:
        raise OfficeEditError("Workbook could not be opened") from exc
    count = 0
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                shown = "" if cell.value is None else str(cell.value)
                if needle not in shown:
                    continue
                current = cell.font
                line = current.underline
                if underline is True:
                    line = "single"
                elif underline is False:
                    line = "none"
                cell.font = Font(
                    name=font or current.name,
                    size=size_pt or current.size,
                    bold=current.bold if bold is None else bold,
                    italic=current.italic if italic is None else italic,
                    underline=line,
                    color=color or current.color,
                )
                if fill:
                    cell.fill = PatternFill("solid", fgColor=fill)
                count += 1
    if count:
        book.save(path)
    return count


def apply_named_style(path: Path, needle: str, style_name: str) -> int:
    """Point matching paragraphs at a named style. docx and odt only."""
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _docx_named_style(path, needle, style_name)
    if suffix == ".odt":
        return _odt_named_style(path, needle, style_name)
    raise OfficeEditError("Named styles are available for docx and odt")


def _docx_named_style(path: Path, needle: str, style_name: str) -> int:
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(str(path))
    count = 0
    try:
        for para in _docx_paragraphs(doc):
            if needle not in para.text:
                continue
            para.style = style_name
            count += 1
    except (KeyError, ValueError) as exc:
        raise OfficeEditError(f"Style {style_name} was not found") from exc
    if count:
        for para in _docx_paragraphs(doc):
            if needle not in para.text:
                continue
            fonts = para.style.font
            if fonts.name:
                for run in para.runs:
                    node = run._element.get_or_add_rPr()
                    family = node.find(qn("w:rFonts"))
                    if family is None:
                        family = ET.SubElement(node, qn("w:rFonts"))
                    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                        family.set(qn(attr), fonts.name)
        doc.save(str(path))
    return count


def _docx_paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def _odt_named_style(path: Path, needle: str, style_name: str) -> int:
    from holix_office.edit import (
        _clone_info,
        _prepare,
        _read_zip,
        _serialize,
        _write_zip,
    )

    members = _read_zip(path)
    updated = []
    total = 0
    found = False
    for info, data in members:
        if info.filename != "content.xml":
            updated.append((info, data))
            continue
        root = _prepare(data)
        styles = root.find(f"{{{_OFFICE}}}automatic-styles")
        if styles is None:
            styles = ET.Element(f"{{{_OFFICE}}}automatic-styles")
            root.insert(0, styles)
        if not any(el.get(f"{{{_STYLE}}}name") == style_name for el in styles):
            _add_odt_style(styles, style_name)
        for para in root.iter():
            if para.tag not in {f"{{{_TEXT}}}p", f"{{{_TEXT}}}h"}:
                continue
            text = "".join(para.itertext())
            if needle not in text:
                continue
            para.set(f"{{{_TEXT}}}style-name", style_name)
            total += 1
        updated.append((_clone_info(info), _serialize(root)))
        found = True
    if not found:
        raise OfficeEditError("content.xml is missing")
    if total:
        _write_zip(path, updated)
    return total


def _add_odt_style(styles: ET.Element, name: str) -> None:
    style = ET.SubElement(styles, f"{{{_STYLE}}}style")
    style.set(f"{{{_STYLE}}}name", name)
    style.set(f"{{{_STYLE}}}family", "paragraph")
    size = _HEADING_SIZE.get(name)
    if size is None:
        return
    props = ET.SubElement(style, f"{{{_STYLE}}}text-properties")
    props.set(f"{{{_FO}}}font-size", f"{size}pt")
    props.set(f"{{{_FO}}}font-weight", "bold")


def _drawing_paragraphs(root: ET.Element) -> list[ET.Element]:
    parents = {child: parent for parent in root.iter() for child in list(parent)}
    found = []
    for el in root.iter():
        if el.tag != f"{{{_A}}}p":
            continue
        cursor = parents.get(el)
        nested = False
        while cursor is not None:
            if cursor.tag == f"{{{_A}}}p":
                nested = True
                break
            cursor = parents.get(cursor)
        if not nested:
            found.append(el)
    return found


def _pptx_text_nodes(para: ET.Element) -> list[tuple[ET.Element, ET.Element]]:
    parents = {child: parent for parent in para.iter() for child in list(parent)}
    nodes = []
    for el in para.iter():
        if el.tag != f"{{{_A}}}t" or not el.text:
            continue
        cursor: ET.Element | None = el
        run = None
        while cursor is not None and cursor is not para:
            if cursor.tag == f"{{{_A}}}r":
                run = cursor
                break
            cursor = parents.get(cursor)
        if run is not None:
            nodes.append((run, el))
    return nodes


def _mask_right(text: str, needle: str, skip: int) -> str:
    chars = list(text)
    hidden = 0
    cursor = len(text)
    while hidden < skip:
        start = text.rfind(needle, 0, cursor)
        if start < 0:
            break
        for index in range(start, start + len(needle)):
            chars[index] = "\x00"
        cursor = start
        hidden += 1
    return "".join(chars)


def _style_span(para, nodes, start, end, **fmt) -> None:
    if not any(fmt.get(key) is not None and fmt.get(key) != "" for key in fmt):
        return
    pos = 0
    overlapping = []
    for run, node in nodes:
        text = node.text or ""
        node_end = pos + len(text)
        if node_end > start and pos < end:
            overlapping.append((run, node, max(0, start - pos), min(len(text), end - pos)))
        pos = node_end
    for run, node, local_start, local_end in reversed(overlapping):
        target = _isolate(para, run, node, local_start, local_end)
        _style_run(target, **fmt)


def _isolate(para, run, node, local_start, local_end):
    text = node.text or ""
    local_start = max(0, min(local_start, len(text)))
    local_end = max(local_start, min(local_end, len(text)))
    if local_start == 0 and local_end == len(text):
        return run
    before, mid, after = text[:local_start], text[local_start:local_end], text[local_end:]
    parent = next((item for item in para.iter() if run in list(item)), None)
    if parent is None or not mid:
        return run
    index = list(parent).index(run)
    node.text = mid
    if before:
        parent.insert(index, _run_like(run, node, before))
        index += 1
    if after:
        parent.insert(index + 1, _run_like(run, node, after))
    return run


def _run_like(run, source, text):
    clone = ET.Element(run.tag, run.attrib)
    props = run.find(f"{{{_A}}}rPr")
    if props is not None:
        clone.append(ET.fromstring(ET.tostring(props)))
    node = ET.Element(source.tag, source.attrib)
    node.text = text
    clone.append(node)
    return clone


def _style_run(run, *, bold, italic, underline, color, font, size_pt) -> None:
    props = run.find(f"{{{_A}}}rPr")
    if props is None:
        props = ET.Element(f"{{{_A}}}rPr")
        run.insert(0, props)
    if bold is True:
        props.set("b", "1")
    elif bold is False:
        props.set("b", "0")
    if italic is True:
        props.set("i", "1")
    elif italic is False:
        props.set("i", "0")
    if underline is True:
        props.set("u", "sng")
    elif underline is False:
        props.set("u", "none")
    if size_pt:
        props.set("sz", str(int(round(size_pt * 100))))
    if font:
        latin = props.find(f"{{{_A}}}latin")
        if latin is None:
            latin = ET.SubElement(props, f"{{{_A}}}latin")
        latin.set("typeface", font)
    if color:
        for child in list(props):
            if child.tag == f"{{{_A}}}solidFill":
                props.remove(child)
        solid = ET.SubElement(props, f"{{{_A}}}solidFill")
        tint = ET.SubElement(solid, f"{{{_A}}}srgbClr")
        tint.set("val", color)


def _shade_paragraph(para: ET.Element, fill: str | None) -> None:
    if not fill:
        return
    props = para.find(f"{{{_A}}}pPr")
    if props is None:
        props = ET.Element(f"{{{_A}}}pPr")
        para.insert(0, props)
    for child in list(props):
        if child.tag == f"{{{_A}}}solidFill":
            props.remove(child)
    solid = ET.SubElement(props, f"{{{_A}}}solidFill")
    tint = ET.SubElement(solid, f"{{{_A}}}srgbClr")
    tint.set("val", fill)
