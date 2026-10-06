"""Add and edit slides in pptx and odp presentations."""

from __future__ import annotations

import xml.etree.ElementTree as ET
from pathlib import Path

from holix_office.edit import OfficeEditError, _require_office_file
from holix_office.odf_xml import (
    DRAW,
    body,
    edit_content,
    pages,
    read_content,
    text_box,
)

_LAYOUTS = {
    "title": 0,
    "title_body": 1,
    "section": 2,
    "two": 3,
    "title_only": 5,
    "blank": 6,
}


def slide_office(
    path: Path,
    action: str,
    *,
    slide: int = 1,
    title: str = "",
    body_text: str = "",
    layout: str = "title_body",
    notes: str = "",
) -> dict[str, object]:
    path = _require_office_file(path)
    suffix = path.suffix.lower()
    if suffix not in {".pptx", ".odp"}:
        raise OfficeEditError("Slides are edited in pptx and odp files")
    verb = (action or "").strip().lower()
    if suffix == ".pptx":
        return _pptx(path, verb, slide=slide, title=title, body_text=body_text, layout=layout, notes=notes)
    return _odp(path, verb, slide=slide, title=title, body_text=body_text, notes=notes)


def _pptx(path, action, *, slide, title, body_text, layout, notes) -> dict[str, object]:
    from pptx import Presentation

    deck = Presentation(str(path))
    if action == "list":
        return {"ok": True, "slides": [_pptx_slide(item, index) for index, item in enumerate(deck.slides, start=1)]}
    if action == "add":
        index = _LAYOUTS.get((layout or "title_body").strip().lower())
        if index is None:
            raise OfficeEditError("layout must be title, title_body, section, two, title_only, or blank")
        made = deck.slides.add_slide(deck.slide_layouts[index])
        _set_placeholders(made, title, body_text)
        deck.save(str(path))
        return {"ok": True, "slide": len(deck.slides), "action": "add"}
    target = _pptx_at(deck, slide)
    if action == "set":
        _set_placeholders(target, title, body_text)
    elif action == "notes":
        if not str(notes).strip():
            raise OfficeEditError("Notes text is empty")
        target.notes_slide.notes_text_frame.text = str(notes)
    elif action == "delete":
        _delete_pptx(deck, slide)
        deck.save(str(path))
        return {"ok": True, "action": "delete", "slides": len(deck.slides)}
    else:
        raise OfficeEditError("action must be list, add, set, notes, or delete")
    deck.save(str(path))
    return {"ok": True, "action": action, "slide": max(1, int(slide or 1))}


def _odp(path, action, *, slide, title, body_text, notes) -> dict[str, object]:
    if action == "list":
        root = read_content(path)
        found = []
        for index, page in enumerate(pages(root), start=1):
            texts = ["".join(el.itertext()).strip() for el in page.iter() if el.tag.endswith("}p")]
            found.append(
                {
                    "slide": index,
                    "name": page.get(f"{{{DRAW}}}name") or f"page{index}",
                    "text": [item for item in texts if item],
                }
            )
        return {"ok": True, "slides": found}
    if action == "add":
        def add(root: ET.Element) -> None:
            page = ET.SubElement(body(root, "presentation"), f"{{{DRAW}}}page")
            page.set(f"{{{DRAW}}}name", f"page{len(pages(root))}")
            if title:
                page.append(text_box(title, "1cm", "0.6cm", "24cm", "2.4cm"))
            if body_text:
                page.append(text_box(body_text, "1cm", "3.4cm", "24cm", "12cm"))

        edit_content(path, add)
        return {"ok": True, "action": "add"}
    if action == "delete":
        removed = {"ok": False}

        def delete(root: ET.Element) -> None:
            found_pages = pages(root)
            index = max(1, int(slide or 1)) - 1
            if index >= len(found_pages):
                raise OfficeEditError("Slide was not found")
            parent = next(item for item in root.iter() if found_pages[index] in list(item))
            parent.remove(found_pages[index])
            removed["ok"] = True

        edit_content(path, delete)
        return {"ok": True, "action": "delete"}
    if action not in {"set", "notes"}:
        raise OfficeEditError("action must be list, add, set, notes, or delete")

    def change(root: ET.Element) -> None:
        found_pages = pages(root)
        index = max(1, int(slide or 1)) - 1
        if index >= len(found_pages):
            raise OfficeEditError("Slide was not found")
        page = found_pages[index]
        if action == "notes":
            if not str(notes).strip():
                raise OfficeEditError("Notes text is empty")
            page.append(text_box(f"Notes: {notes}", "1cm", "16cm", "24cm", "2cm"))
            return
        paragraphs = [el for el in page.iter() if el.tag.endswith("}p")]
        lines = [title, body_text]
        if not paragraphs:
            if title:
                page.append(text_box(title, "1cm", "0.6cm", "24cm", "2.4cm"))
            if body_text:
                page.append(text_box(body_text, "1cm", "3.4cm", "24cm", "12cm"))
            return
        for paragraph, line in zip(paragraphs, [item for item in lines if item], strict=False):
            paragraph.text = line
            for child in list(paragraph):
                paragraph.remove(child)

    edit_content(path, change)
    return {"ok": True, "action": action, "slide": max(1, int(slide or 1))}


def _pptx_slide(slide, index: int) -> dict[str, object]:
    texts = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            text = shape.text_frame.text.strip()
            if text:
                texts.append(text)
    notes = ""
    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame.text.strip()
    return {"slide": index, "text": texts, "notes": notes}


def _set_placeholders(slide, title: str, body_text: str) -> None:
    if title and getattr(slide.shapes, "title", None) is not None and slide.shapes.title is not None:
        slide.shapes.title.text = title
    if not body_text:
        return
    for shape in slide.placeholders:
        if shape.placeholder_format.idx == 1:
            shape.text = body_text
            return
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and shape != slide.shapes.title:
            shape.text_frame.text = body_text
            return


def _pptx_at(deck, slide: int):
    if len(deck.slides) == 0:
        raise OfficeEditError("Presentation has no slides")
    index = max(1, int(slide or 1)) - 1
    if index >= len(deck.slides):
        raise OfficeEditError("Slide was not found")
    return deck.slides[index]


def _delete_pptx(deck, slide: int) -> None:
    index = max(1, int(slide or 1)) - 1
    id_list = deck.slides._sldIdLst
    if index >= len(id_list):
        raise OfficeEditError("Slide was not found")
    item = id_list[index]
    rel = item.get(
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    )
    if rel:
        deck.part.drop_rel(rel)
    id_list.remove(item)
